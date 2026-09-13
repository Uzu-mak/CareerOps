from __future__ import annotations

from io import BytesIO
from pathlib import Path
from datetime import datetime, timezone
import re
from pypdf import PdfReader
from docx import Document
from fastapi import UploadFile
from sqlalchemy.orm import Session

from careerops.core.config import settings
from careerops.db.models import CandidateProfile

ALLOWED_EXTENSIONS = {'.pdf', '.docx', '.txt'}
MAX_RESUME_BYTES = 8 * 1024 * 1024


def _clean_text(text: str) -> str:
    text = text.replace('\x00', ' ')
    lines = [re.sub(r'\s+', ' ', x).strip() for x in text.splitlines()]
    return '\n'.join(x for x in lines if x)


def extract_resume_text(data: bytes, suffix: str) -> str:
    suffix = suffix.lower()
    if suffix == '.pdf':
        reader = PdfReader(BytesIO(data))
        return _clean_text('\n'.join((page.extract_text() or '') for page in reader.pages))
    if suffix == '.docx':
        doc = Document(BytesIO(data))
        chunks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(' | '.join(cell.text for cell in row.cells))
        return _clean_text('\n'.join(chunks))
    if suffix == '.txt':
        return _clean_text(data.decode('utf-8', errors='replace'))
    raise ValueError('Unsupported resume file type. Use PDF, DOCX, or TXT.')


async def save_baseline_resume(db: Session, profile: CandidateProfile, upload: UploadFile) -> dict:
    original = Path(upload.filename or 'resume').name
    suffix = Path(original).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError('Unsupported resume file type. Use PDF, DOCX, or TXT.')
    data = await upload.read()
    if not data:
        raise ValueError('The uploaded resume is empty.')
    if len(data) > MAX_RESUME_BYTES:
        raise ValueError('Resume is too large. Maximum size is 8 MB.')
    text = extract_resume_text(data, suffix)
    if len(text) < 120:
        raise ValueError('CareerOps could not extract enough text from this resume. Try a text-based PDF or DOCX.')

    stamp = datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')
    path = settings.uploads_path / f'master_resume_{profile.id}_{stamp}{suffix}'
    path.write_bytes(data)
    profile.baseline_resume_filename = original
    profile.baseline_resume_path = str(path)
    profile.baseline_resume_text = text
    profile.baseline_resume_uploaded_at = datetime.now(timezone.utc)
    db.add(profile)
    db.commit()
    db.refresh(profile)
    return {
        'filename': original,
        'path': str(path),
        'characters_extracted': len(text),
        'uploaded_at': profile.baseline_resume_uploaded_at,
    }


def baseline_bullets(profile: CandidateProfile) -> list[str]:
    """Recover bullet-like claims from the master resume, joining wrapped PDF lines."""
    raw = (profile.baseline_resume_text or '').splitlines()
    bullets: list[str] = []
    current: list[str] = []
    section_headers = {'SUMMARY','EDUCATION','EXPERIENCE','PROJECTS','TECHNICAL SKILLS','SKILLS','CERTIFICATIONS'}
    for raw_line in raw:
        line = raw_line.strip()
        if not line:
            continue
        is_bullet = line.startswith(('•','●','▪','- '))
        clean = line.lstrip('•●▪- ').strip()
        if is_bullet:
            if current:
                bullets.append(' '.join(current))
            current = [clean]
            continue
        if current:
            # PDF extraction commonly wraps one bullet across several physical lines.
            if line.upper() in section_headers or (' | ' in line and len(line) < 180 and not line.endswith(('.', ';', ','))):
                bullets.append(' '.join(current)); current=[]
            else:
                current.append(clean)
    if current:
        bullets.append(' '.join(current))
    return [re.sub(r'\s+', ' ', b).strip() for b in bullets if len(b.strip()) >= 30]


def baseline_lines(profile: CandidateProfile) -> list[str]:
    bullets = baseline_bullets(profile)
    if bullets:
        return bullets
    # DOCX/plain-text fallback: keep sentence-like accomplishment lines, excluding headers/contact/skills.
    out=[]
    for x in (profile.baseline_resume_text or '').splitlines():
        line=x.strip()
        low=line.lower()
        if len(line) < 35 or '@' in line or 'linkedin' in low or 'github' in low:
            continue
        if re.match(r'^(languages|full stack|applied ai|data & systems|devops|collaboration|coursework):', low):
            continue
        if line.isupper():
            continue
        out.append(line)
    return out


def select_baseline_lines(profile: CandidateProfile, keywords: list[str], limit: int = 18) -> list[str]:
    """Select job-relevant, user-authored claims from the uploaded master resume."""
    lines = baseline_lines(profile)
    if not lines:
        return []
    keys = [k.lower() for k in keywords if k]
    scored = []
    for i, line in enumerate(lines):
        low = line.lower()
        score = sum(3 for k in keys if k in low)
        if any(token in low for token in ['engineer', 'project', 'python', 'sql', 'api', 'ai', 'data', 'software']):
            score += 1
        scored.append((score, -i, line))
    scored.sort(reverse=True)
    selected = [line for score, _, line in scored if score > 0][:limit]
    if len(selected) < min(6, len(lines)):
        seen = set(selected)
        for line in lines:
            if line not in seen:
                selected.append(line); seen.add(line)
            if len(selected) >= min(limit, max(6, len(lines))):
                break
    return selected[:limit]


def extract_section_lines(profile: CandidateProfile, section: str, stop_sections: tuple[str, ...] = ('SUMMARY','EDUCATION','EXPERIENCE','PROJECTS','TECHNICAL SKILLS','SKILLS','CERTIFICATIONS')) -> list[str]:
    lines=[x.strip() for x in (profile.baseline_resume_text or '').splitlines() if x.strip()]
    target=section.upper()
    start=None
    for i,line in enumerate(lines):
        if line.upper()==target:
            start=i+1; break
    if start is None: return []
    out=[]
    stops={x.upper() for x in stop_sections if x.upper()!=target}
    for line in lines[start:]:
        if line.upper() in stops: break
        out.append(line)
    return out
